"""运行时工作台参考资料的唯一读取入口。

参考资料从 ``docs/`` 迁移到模型服务后，模型接口只通过本模块读取固定的
``data/reference`` 目录。调用方不能传入路径，也不能覆盖目录位置。
"""

from __future__ import annotations

import math
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


REFERENCE_DIR = Path(__file__).resolve().parents[1] / "data" / "reference"

DATASET_FILES: dict[str, str] = {
    "cost_data": "cost_data.xlsx",
    "price_elasticity": "price_elasticity.xlsx",
    "strategy_library": "strategy_library.xlsx",
}

REQUIRED_COLUMNS: dict[str, tuple[str, ...]] = {
    "cost_data": ("品类", "型号", "建议零售价", "成本价"),
    "price_elasticity": (
        "品类",
        "系列",
        "型号",
        "均价",
        "均销",
        "价格弹性系数",
        "波动分类",
        "弹性分类",
        "产品分类",
    ),
}

KNOWLEDGE_FILES = (
    "strategy_library.xlsx",
    "sales_factor_knowledge.md",
    "architecture.md",
)


class ReferenceDataError(ValueError):
    """参考数据不存在、格式非法或缺少约定列。"""


def _clean_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        return value.isoformat()
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except Exception:  # noqa: BLE001
            return str(value)
    if hasattr(value, "item"):
        try:
            return value.item()
        except Exception:  # noqa: BLE001
            pass
    return value


def _header_name(value: Any) -> str:
    cleaned = _clean_value(value)
    return "" if cleaned is None else str(cleaned).strip()


def _source_for(path: Path) -> str:
    """返回稳定的仓库相对来源，不把绝对机器路径暴露给客户端。"""
    # The model is mounted at different roots in local and container runs
    # (for example ``services/icewash-model`` vs ``/app``).  Derive the
    # public source from the fixed reference directory instead of assuming a
    # particular number of parent directories.
    try:
        relative = path.resolve().relative_to(REFERENCE_DIR.resolve())
    except ValueError:
        return path.name
    return Path("services/icewash-model/data/reference", relative).as_posix()


def _dataset_path(dataset: str) -> Path:
    try:
        filename = DATASET_FILES[dataset]
    except KeyError as exc:
        raise ReferenceDataError(
            f"不支持的参考数据集: {dataset}；仅允许 {', '.join(DATASET_FILES)}"
        ) from exc
    return REFERENCE_DIR / filename


def _validate_required_columns(dataset: str, filename: str, columns: list[str]) -> None:
    required = REQUIRED_COLUMNS[dataset]
    present = set(columns)
    missing = [column for column in required if column not in present]
    if missing:
        raise ReferenceDataError(f"文件 {filename} 缺少列: {', '.join(missing)}")


def _normalize_frame(frame: pd.DataFrame, *, dataset: str, filename: str) -> tuple[list[str], list[dict[str, Any]]]:
    columns = [_header_name(column) for column in frame.columns]
    _validate_required_columns(dataset, filename, columns)
    # 只对列名做去空格归一，保留导出文件的全部列和值。
    frame = frame.copy()
    frame.columns = columns
    rows = [
        {column: _clean_value(value) for column, value in row.items()}
        for row in frame.to_dict(orient="records")
    ]
    return columns, rows


def _read_excel_frame(path_or_buffer: Any, *, dataset: str, filename: str) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        frame = pd.read_excel(path_or_buffer, sheet_name=0)
    except Exception as exc:  # noqa: BLE001
        raise ReferenceDataError(f"文件 {filename} 无法读取为 Excel: {exc}") from exc
    return _normalize_frame(frame, dataset=dataset, filename=filename)


def _read_strategy(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    filename = path.name
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ReferenceDataError(f"文件 {filename} 无法读取为 Excel: {exc}") from exc
    try:
        if "Sheet1" not in workbook.sheetnames:
            raise ReferenceDataError(f"文件 {filename} 缺少工作表 Sheet1")
        sheet = workbook["Sheet1"]
        header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [_header_name(value) for value in header_values]
        if not headers or any(not header for header in headers):
            raise ReferenceDataError(f"文件 {filename} 存在空表头")
        duplicates = sorted({header for header in headers if headers.count(header) > 1})
        if duplicates:
            raise ReferenceDataError(f"文件 {filename} 存在重复表头: {', '.join(duplicates)}")

        rows: list[dict[str, Any]] = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            normalized = [_clean_value(value) for value in values[: len(headers)]]
            if not any(value is not None and str(value).strip() for value in normalized):
                continue
            normalized.extend([None] * (len(headers) - len(normalized)))
            rows.append(dict(zip(headers, normalized, strict=True)))
        if not rows:
            raise ReferenceDataError(f"文件 {filename} 没有数据行")
        return headers, rows
    finally:
        workbook.close()


def load_workbench_dataset(dataset: str) -> dict[str, Any]:
    """读取允许的工作台数据集并输出 JSON-safe 规范化结果。"""
    path = _dataset_path(dataset)
    if not path.is_file():
        raise ReferenceDataError(f"参考文件不存在: {path.name}")
    if dataset == "strategy_library":
        columns, rows = _read_strategy(path)
    else:
        columns, rows = _read_excel_frame(path, dataset=dataset, filename=path.name)
    return {
        "dataset": dataset,
        "source": _source_for(path),
        "columns": columns,
        "rows": rows,
        "row_count": len(rows),
    }


def parse_cost_upload(content: bytes, filename: str) -> tuple[list[str], list[dict[str, Any]], pd.DataFrame]:
    """校验成本上传并返回规范化表格，调用方负责原子落盘。"""
    lower_name = (filename or "").lower()
    if not lower_name.endswith((".csv", ".xlsx")):
        raise ReferenceDataError("成本文件仅支持 .csv 或 .xlsx")
    if not content:
        raise ReferenceDataError("上传文件为空")
    buffer = BytesIO(content)
    try:
        if lower_name.endswith(".csv"):
            frame = pd.read_csv(buffer)
        else:
            frame = pd.read_excel(buffer, sheet_name=0)
    except Exception as exc:  # noqa: BLE001
        raise ReferenceDataError(f"文件 {filename} 无法读取: {exc}") from exc
    if frame.empty:
        raise ReferenceDataError(f"文件 {filename} 没有数据行")
    columns, rows = _normalize_frame(frame, dataset="cost_data", filename=filename)
    frame = frame.copy()
    frame.columns = columns
    return columns, rows, frame


def knowledge_markdown() -> dict[str, str]:
    """按固定章节顺序拼接知识库；策略 Excel 不作为可执行代码读取。"""
    sales_path = REFERENCE_DIR / "knowledge" / "sales_factor_knowledge.md"
    architecture_path = REFERENCE_DIR / "knowledge" / "architecture.md"
    missing = [path.name for path in (sales_path, architecture_path) if not path.is_file()]
    if missing:
        raise ReferenceDataError(f"知识库文件不存在: {', '.join(missing)}")
    sales = sales_path.read_text(encoding="utf-8")
    architecture = architecture_path.read_text(encoding="utf-8")
    markdown = (
        "# 模拟策略库\n\n"
        "策略目录来源于 `strategy_library.xlsx`，可执行策略和公式由模型 whatif 代码提供。\n\n"
        "# 量化知识库\n\n"
        f"{sales.rstrip()}\n\n"
        "# 推导指南\n\n"
        f"{architecture.rstrip()}\n"
    )
    return {
        "title": "策略知识库",
        "source": "; ".join(KNOWLEDGE_FILES),
        "markdown": markdown,
    }
